# Copyright (c) 2026, Healthcare contributors
"""Fill missing Healthcare Practitioner names from an uploaded practitioner list.

Some legacy practitioner records were created from Oracle appointment codes and
only carry the code (``doctors_id``) with no full name. The clinic practitioner
list keeps the full employee name next to the employee code and the user (login)
code, so it can be used to backfill those records.

Rows are matched on the legacy code (``doctors_id`` / document name) against, in
order, *Employee Code*, *User Code* and finally *Login Name* (only possible for
practitioners that have a linked User). Codes that map to more than one name are
skipped and reported instead of guessing.

``preview_practitioner_name_import`` reports what would change and
``run_practitioner_name_import`` writes ``first_name``, ``last_name`` and
``practitioner_name`` on practitioners whose full name is still empty or is just
the legacy code.
"""

from __future__ import annotations

from typing import Any

import frappe
from frappe import _

from healthcare.api.patient_info_import import (
	_cell_text,
	_clean_oracle_num,
	_excel_file_path,
	_require_admin,
)

EXCEL_HEADER_MAP = {
	"EMPLOYEE CODE": "employee_code",
	"EMPLOYEE NAME": "full_name",
	"STAFF CODE": "employee_code",
	"STAFF NAME": "full_name",
	"USER CODE": "user_code",
	"LOGIN NAME": "login_name",
	"FULL NAME": "full_name",
	"NAME": "full_name",
}

MAX_REPORTED_CODES = 50


def _normalize_header(cell: Any) -> str:
	if cell is None:
		return ""
	text = " ".join(str(cell).strip().upper().replace("_", " ").split())
	return EXCEL_HEADER_MAP.get(text, text.lower().replace(" ", "_"))


def _clean_full_name(value: Any) -> str:
	return " ".join(_cell_text(value).split())


def _sort_codes(codes: list[str] | set[str]) -> list[str]:
	"""Sort legacy codes numerically first, then any non-numeric leftovers."""

	def key(code: str) -> tuple[int, Any]:
		return (0, int(code)) if code.isdigit() else (1, code)

	return sorted(codes, key=key)


def _parse_sheet_rows(ws) -> list[dict[str, str]]:
	rows_iter = ws.iter_rows(values_only=True)
	try:
		header_row = next(rows_iter)
	except StopIteration:
		return []

	headers = [_normalize_header(h) for h in header_row]
	if "employee_code" not in headers or "full_name" not in headers:
		return []

	parsed: list[dict[str, str]] = []
	for raw in rows_iter:
		if not raw or all(cell is None or str(cell).strip() == "" for cell in raw):
			continue
		row: dict[str, Any] = {}
		for idx, key in enumerate(headers):
			if not key or idx >= len(raw):
				continue
			row[key] = raw[idx]
		employee_code = _clean_oracle_num(row.get("employee_code"))
		full_name = _clean_full_name(row.get("full_name"))
		if not employee_code or not full_name:
			continue
		parsed.append(
			{
				"employee_code": employee_code,
				"full_name": full_name,
				"user_code": _clean_oracle_num(row.get("user_code")),
				"login_name": _cell_text(row.get("login_name")),
			}
		)
	return parsed


def _parse_excel_rows(file_url: str) -> tuple[list[dict[str, str]], dict[str, int]]:
	try:
		import openpyxl
	except ImportError:
		frappe.throw(
			_(
				"openpyxl is required to read Excel files. "
				"Install it in the bench environment: pip install openpyxl"
			)
		)

	path = _excel_file_path(file_url)
	wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
	by_employee_code: dict[str, dict[str, str]] = {}
	sheet_row_counts: dict[str, int] = {}
	try:
		for sheet_name in wb.sheetnames:
			sheet_rows = _parse_sheet_rows(wb[sheet_name])
			sheet_row_counts[sheet_name] = len(sheet_rows)
			for row in sheet_rows:
				code = row["employee_code"]
				existing = by_employee_code.get(code)
				if not existing:
					by_employee_code[code] = dict(row)
					continue
				# Re-engaged staff appear more than once; keep the fullest name but
				# do not lose a user code that only one of the rows carries.
				if len(row["full_name"]) > len(existing["full_name"]):
					existing["full_name"] = row["full_name"]
				existing["user_code"] = existing.get("user_code") or row.get("user_code") or ""
				existing["login_name"] = existing.get("login_name") or row.get("login_name") or ""
	finally:
		wb.close()
	return list(by_employee_code.values()), sheet_row_counts


def _index_names_by_code(
	rows: list[dict[str, str]],
) -> tuple[dict[str, str], dict[str, set[str]], dict[str, set[str]]]:
	by_employee_code: dict[str, str] = {}
	by_user_code: dict[str, set[str]] = {}
	by_login_name: dict[str, set[str]] = {}
	for row in rows:
		by_employee_code[row["employee_code"]] = row["full_name"]
		user_code = (row.get("user_code") or "").strip()
		if user_code:
			by_user_code.setdefault(user_code, set()).add(row["full_name"])
		login_name = (row.get("login_name") or "").strip().upper()
		if login_name:
			by_login_name.setdefault(login_name, set()).add(row["full_name"])
	return by_employee_code, by_user_code, by_login_name


def _login_keys_by_user(user_ids: list[str]) -> dict[str, set[str]]:
	"""Map each linked User to the uppercase keys its login name can match."""
	if not user_ids:
		return {}

	keys_by_user: dict[str, set[str]] = {}
	for user in frappe.get_all(
		"User", filters={"name": ("in", user_ids)}, fields=["name", "username"]
	):
		values = {user.name, user.username or "", (user.name or "").split("@")[0]}
		keys_by_user[user.name] = {value.strip().upper() for value in values if value.strip()}
	return keys_by_user


def _practitioners_without_full_name() -> list[dict]:
	"""Practitioners whose full name is empty or still only the legacy code."""
	return frappe.db.sql(
		"""
		select name, doctors_id, first_name, last_name, practitioner_name, user_id
		from `tabHealthcare Practitioner`
		where coalesce(trim(practitioner_name), '') = ''
			or trim(practitioner_name) = trim(coalesce(doctors_id, ''))
			or trim(practitioner_name) = trim(name)
			or trim(practitioner_name) regexp '^[0-9]+$'
		order by name
		""",
		as_dict=True,
	)


def _candidate_codes(practitioner: dict) -> list[str]:
	codes: list[str] = []
	for value in (practitioner.get("doctors_id"), practitioner.get("name")):
		code = _clean_oracle_num(value)
		if code and code not in codes:
			codes.append(code)
	return codes


def _resolve_full_name(
	practitioner: dict,
	by_employee_code: dict[str, str],
	by_user_code: dict[str, set[str]],
	by_login_name: dict[str, set[str]] | None = None,
	login_keys_by_user: dict[str, set[str]] | None = None,
) -> tuple[str | None, str | None]:
	"""Return (full name, source). Source is None when nothing matched."""
	codes = _candidate_codes(practitioner)

	employee_names = {by_employee_code[code] for code in codes if code in by_employee_code}
	if len(employee_names) == 1:
		return employee_names.pop(), "employee_code"
	if len(employee_names) > 1:
		return None, "ambiguous"

	user_names: set[str] = set()
	for code in codes:
		user_names.update(by_user_code.get(code, set()))
	if len(user_names) == 1:
		return user_names.pop(), "user_code"
	if len(user_names) > 1:
		return None, "ambiguous"

	login_names: set[str] = set()
	for key in (login_keys_by_user or {}).get(practitioner.get("user_id"), set()):
		login_names.update((by_login_name or {}).get(key, set()))
	if len(login_names) == 1:
		return login_names.pop(), "login_name"
	if len(login_names) > 1:
		return None, "ambiguous"

	return None, None


def _name_fields(full_name: str) -> dict[str, str]:
	parts = full_name.split()
	return {
		"first_name": parts[0],
		"last_name": " ".join(parts[1:]),
		"practitioner_name": full_name,
	}


def _build_update_plan(rows: list[dict[str, str]]) -> dict[str, Any]:
	by_employee_code, by_user_code, by_login_name = _index_names_by_code(rows)
	candidates = _practitioners_without_full_name()
	login_keys_by_user = _login_keys_by_user(
		[practitioner["user_id"] for practitioner in candidates if practitioner.get("user_id")]
	)

	plan: list[dict[str, str]] = []
	unmatched: list[str] = []
	ambiguous: list[str] = []
	matched_by = {"employee_code": 0, "user_code": 0, "login_name": 0}

	for practitioner in candidates:
		full_name, source = _resolve_full_name(
			practitioner, by_employee_code, by_user_code, by_login_name, login_keys_by_user
		)
		if not full_name:
			if source == "ambiguous":
				ambiguous.append(practitioner["name"])
			else:
				unmatched.append(practitioner["name"])
			continue
		matched_by[source] += 1
		plan.append(
			{
				"practitioner": practitioner["name"],
				"doctors_id": practitioner.get("doctors_id") or practitioner["name"],
				"source": source,
				**_name_fields(full_name),
			}
		)

	return {
		"plan": plan,
		"unmatched": _sort_codes(unmatched),
		"ambiguous": _sort_codes(ambiguous),
		"matched_by": matched_by,
	}


def _require_file_url(file_url: str) -> None:
	if not (file_url or "").strip():
		frappe.throw(_("Please upload the practitioner list Excel file."))


@frappe.whitelist()
def preview_practitioner_name_import(file_url: str) -> dict:
	_require_admin()
	_require_file_url(file_url)

	rows, sheet_row_counts = _parse_excel_rows(file_url)
	if not rows:
		frappe.throw(
			_(
				"No names found in the uploaded file. "
				"The sheet needs 'Employee Code' and 'Name' columns."
			)
		)

	built = _build_update_plan(rows)
	plan = built["plan"]
	return {
		"excel_rows": len(rows),
		"sheets": sheet_row_counts,
		"updatable": len(plan),
		"matched_by_employee_code": built["matched_by"]["employee_code"],
		"matched_by_user_code": built["matched_by"]["user_code"],
		"matched_by_login_name": built["matched_by"]["login_name"],
		"unmatched": built["unmatched"][:MAX_REPORTED_CODES],
		"unmatched_count": len(built["unmatched"]),
		"ambiguous": built["ambiguous"][:MAX_REPORTED_CODES],
		"ambiguous_count": len(built["ambiguous"]),
		"sample_updates": plan[:10],
	}


@frappe.whitelist()
def run_practitioner_name_import(file_url: str) -> dict:
	_require_admin()
	_require_file_url(file_url)

	rows, _sheet_row_counts = _parse_excel_rows(file_url)
	built = _build_update_plan(rows)
	plan = built["plan"]

	for item in plan:
		frappe.db.set_value(
			"Healthcare Practitioner",
			item["practitioner"],
			{
				"first_name": item["first_name"],
				"last_name": item["last_name"],
				"practitioner_name": item["practitioner_name"],
			},
		)

	if plan:
		frappe.db.commit()
		frappe.clear_cache(doctype="Healthcare Practitioner")

	return {
		"ok": True,
		"updated": len(plan),
		"matched_by_employee_code": built["matched_by"]["employee_code"],
		"matched_by_user_code": built["matched_by"]["user_code"],
		"matched_by_login_name": built["matched_by"]["login_name"],
		"unmatched": built["unmatched"][:MAX_REPORTED_CODES],
		"unmatched_count": len(built["unmatched"]),
		"ambiguous": built["ambiguous"][:MAX_REPORTED_CODES],
		"ambiguous_count": len(built["ambiguous"]),
	}
